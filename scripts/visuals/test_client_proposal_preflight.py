"""Neutral fixtures verify missing-brand and internal-copy failures."""
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from PIL import Image
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptInches
from client_proposal_preflight import inspect

class ClientPreflightTests(unittest.TestCase):
 def setUp(self):
  self.temp=TemporaryDirectory();self.root=Path(self.temp.name)
  self.assets=[]
  for n,c in [('brand','blue'),('client','green'),('cover','white')]:
   f=self.root/(n+'.png');Image.new('RGB',(100,30),c).save(f);self.assets.append(f)
 def tearDown(self):self.temp.cleanup()
 def word(self,header=True,client=True,cover=True,copy='About Mindmap Digital'):
  d=Document();d.add_heading('A business proposal',0);d.add_picture(str(self.assets[0]),width=Inches(1))
  if client:d.add_picture(str(self.assets[1]),width=Inches(1))
  if cover:d.add_picture(str(self.assets[2]),width=Inches(3))
  d.add_page_break();d.add_heading(copy,1)
  if header:d.sections[0].header.paragraphs[0].add_run().add_picture(str(self.assets[0]),width=Inches(1))
  f=self.root/'proposal.docx';d.save(f);return f
 def check_word(self,**kw):return inspect(self.word(**kw),None,*self.assets[:2])
 def test_branded_package_passes(self):self.assertTrue(self.check_word()['passed'])
 def test_header_is_required(self):self.assertIn('header missing',' '.join(self.check_word(header=False)['errors']))
 def test_client_logo_is_required(self):self.assertIn('client logo missing',' '.join(self.check_word(client=False)['errors']))
 def test_cover_cannot_be_just_logos(self):self.assertIn('cover graphic missing',' '.join(self.check_word(cover=False)['errors']))
 def test_internal_copy_is_rejected(self):
  r=self.check_word(copy='About Mindmap Digital. The notes indicate a review draft.');self.assertFalse(r['passed']);self.assertIn('internal authoring language',' '.join(r['errors']))
 def test_slide_branding_is_checked_beyond_cover(self):
  p=Presentation();s=p.slides.add_slide(p.slide_layouts[6])
  for i,f in enumerate(self.assets):s.shapes.add_picture(str(f),PptInches(i),0,width=PptInches(1))
  s.shapes.add_textbox(0,PptInches(1),PptInches(5),PptInches(1)).text='About Mindmap Digital'
  p.slides.add_slide(p.slide_layouts[6]);f=self.root/'deck.pptx';p.save(f)
  r=inspect(None,f,*self.assets[:2]);self.assertIn('PowerPoint slide 2: Mindmap logo missing',r['errors'])

if __name__=='__main__':unittest.main()
