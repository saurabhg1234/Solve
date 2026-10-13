"""Local synthetic/source fixtures. Never create a client proposal or approval."""
import hashlib, json, sys, tempfile, unittest, uuid
from pathlib import Path
from zipfile import ZipFile
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
from prepare_sources import prepare
from build_documents import build, inspect
from build_rfi import build as build_rfi
from docx import Document
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont
import pymupdf

REPO=Path(__file__).resolve().parents[1]
def digest(data): return hashlib.sha256(data).hexdigest()

class NativeDocuments(unittest.TestCase):
    def setUp(self):
        self.directory=tempfile.TemporaryDirectory(prefix='zora-native-');self.root=Path(self.directory.name)
    def tearDown(self): self.directory.cleanup()
    def source(self,name,data,mime):
        identity=str(uuid.uuid4());path=self.root/(identity+'.source');path.write_bytes(data)
        item={'id':identity,'filename':name,'mime_type':mime,'sha256':digest(data),'path':path.name}
        (self.root/'sources.json').write_text(json.dumps([item]));return item
    def test_transcript_locations_hashes_and_no_automatic_verification(self):
        item=self.source('call.txt',b'00:01:12 Client: Please confirm receipt.\n00:02:10 Client: No pricing is agreed.','text/plain')
        result=prepare(self.root);self.assertEqual(result['status'],'prepared')
        evidence=json.loads((self.root/'prepared/evidence.json').read_text())['evidence'];self.assertEqual(evidence[0]['timestamp'],'00:01:12');self.assertEqual(evidence[0]['sourceSha256'],item['sha256']);self.assertFalse(evidence[0]['verified'])
        (self.root/item['path']).write_text('changed');self.assertRaises(ValueError,prepare,self.root)
    def test_ocr_batch_resumes_from_hash_bound_cache(self):
        image=Image.new('RGB',(1600,500),'white');draw=ImageDraw.Draw(image);font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',42);draw.text((60,160),'Synthetic OCR source. No client proposal.',fill='black',font=font)
        image_path=self.root/'scan.png';image.save(image_path)
        pdf=pymupdf.open()
        for _ in range(2):page=pdf.new_page(width=800,height=250);page.insert_image(page.rect,filename=str(image_path))
        self.source('scanned.pdf',pdf.tobytes(),'application/pdf');pdf.close()
        self.assertEqual(prepare(self.root,max_ocr_pages=1)['status'],'blocked')
        self.assertEqual(prepare(self.root,max_ocr_pages=1)['status'],'prepared')
        report=json.loads((self.root/'prepared/evidence.json').read_text());self.assertEqual({r['page'] for r in report['evidence']},{1,2});self.assertTrue(all(r['needsVisualVerification'] for r in report['evidence']))
        self.assertEqual(prepare(self.root,max_ocr_pages=0)['status'],'prepared')
    def test_word_xml_and_audio_holds(self):
        template=self.root/'fixture.docx';doc=Document();doc.add_paragraph('Synthetic source paragraph');doc.add_table(rows=1,cols=1).cell(0,0).text='Source table cell';doc.save(template)
        self.source('source.docx',template.read_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document');self.assertEqual(prepare(self.root)['status'],'prepared');records=json.loads((self.root/'prepared/evidence.json').read_text())['evidence'];self.assertTrue(any(r['kind']=='table_cell' for r in records))
        self.source('call.mp3',b'synthetic unsupported audio fixture','audio/mpeg');self.assertEqual(prepare(self.root)['status'],'blocked')
    def word_plan(self):
        path=self.root/'synthetic-approved-template.docx';doc=Document();doc.add_paragraph('LOCAL VERIFICATION — not a client template');p=doc.add_paragraph();r=p.add_run('CLIENT_PLACEHOLDER');r.bold=True;doc.save(path)
        details=inspect(path,'docx');parts=[]
        for part in details['parts']:
            parts.append({'part':part['part'],'edits':[{'index':x['index'],'expected':x['text'],'text':'Synthetic Client'} for x in part['texts'] if x['text']=='CLIENT_PLACEHOLDER'],'retained':[{'index':x['index'],'expected':x['text']} for x in part['texts'] if x['text']!='CLIENT_PLACEHOLDER']})
        return {'format':'docx','template':str(path),'templateSha256':details['templateSha256'],'approvedTemplate':True,'output':'outputs/client/verification.docx','parts':parts}
    def test_word_styles_pdf_and_approval_guards(self):
        plan=self.word_plan();plan['pdf']=True;report=build(plan,self.root);output=self.root/report['output'];self.assertEqual(Document(output).paragraphs[1].text,'Synthetic Client');self.assertTrue(Document(output).paragraphs[1].runs[0].bold)
        with ZipFile(plan['template']) as before,ZipFile(output) as after:self.assertEqual(before.read('word/styles.xml'),after.read('word/styles.xml'));self.assertNotIn('docProps/core.xml',after.namelist())
        with pymupdf.open(self.root/report['pdf']['output']) as pdf:self.assertEqual(len(pdf),1);self.assertIn('Synthetic Client',pdf[0].get_text())
        self.assertRaises(ValueError,build,{**plan,'approvedTemplate':False},self.root)
        self.assertRaises(ValueError,build,{**plan,'output':'../../escape.docx'},self.root)
        self.assertRaises(ValueError,build,{**plan,'templateSha256':'0'*64},self.root)
    def test_ppt_duplicates_keep_styles_and_remove_unselected_slides(self):
        template=REPO/'AI_Solutioning_Agent_BRD_v1.0/02_Templates/Reference deck/Reference_Deck_Slides.pptx';details=inspect(template,'pptx');part=next(p for p in details['parts'] if p['part']=='ppt/slides/slide8.xml')
        slide={'sourceSlide':8,'edits':[{'index':x['index'],'expected':x['text'],'text':'LOCAL VERIFICATION' if x['index']==0 else 'Synthetic fixture only'} for x in part['texts']],'retained':[]}
        plan={'format':'pptx','template':str(template),'templateSha256':details['templateSha256'],'approvedTemplate':True,'output':'outputs/client/verification.pptx','slides':[slide,slide],'pdf':True};report=build(plan,self.root);output=self.root/report['output'];presentation=Presentation(output);self.assertEqual(len(presentation.slides),2)
        with ZipFile(template) as before,ZipFile(output) as after:
            self.assertEqual(before.read('ppt/theme/theme1.xml'),after.read('ppt/theme/theme1.xml'));self.assertFalse(any(n.startswith('ppt/notesSlides/') for n in after.namelist()));self.assertNotIn('ppt/slides/slide1.xml',after.namelist());self.assertNotIn('docProps/thumbnail.jpeg',after.namelist())
        with pymupdf.open(self.root/report['pdf']['output']) as pdf:self.assertEqual(len(pdf),2);self.assertIn('LOCAL VERIFICATION',pdf[0].get_text())
        self.assertRaises(ValueError,build,{**plan,'slides':[{'sourceSlide':999}]},self.root)
        self.assertRaises(ValueError,build,{**plan,'slides':[{'sourceSlide':8,'edits':[],'retained':[]}]},self.root)
    def test_all_five_supplied_templates_compose_without_recreating_layouts(self):
        templates=sorted((REPO/'AI_Solutioning_Agent_BRD_v1.0/02_Templates/Slide templates').glob('*.pptx'))
        approved=[];slides=[]
        for i,template in enumerate(templates):
            details=inspect(template,'pptx');key='base' if i==0 else f'template{i}'
            spec={'template':str(template),'templateSha256':details['templateSha256'],'approvedTemplate':True}
            approved.append({**spec,'key':key})
            part=details['parts'][0];slides.append({'templateKey':key,'sourceSlide':1,'edits':[],'retained':[{'index':x['index'],'expected':x['text']} for x in part['texts']]})
        plan={**approved[0],'format':'pptx','templates':approved[1:],'slides':slides,'pdf':True,'output':'outputs/internal/five-template-verification.pptx'}
        report=build(plan,self.root);self.assertEqual(len(Presentation(self.root/report['output']).slides),5);self.assertEqual(len(report['templateHashes']),5)
        with ZipFile(self.root/report['output']) as generated:
            for i,template in enumerate(templates[1:],1):
                with ZipFile(template) as original:
                    self.assertEqual(original.read('ppt/theme/theme1.xml'),generated.read(f'ppt/zoraTemplate{i}/theme/theme1.xml'))
            self.assertFalse(any('notesSlides/' in n or n.startswith('docProps/') for n in generated.namelist()))
        with pymupdf.open(self.root/report['pdf']['output']) as pdf:self.assertEqual(len(pdf),5);self.assertTrue(all(page.get_text().strip() for page in pdf))
        self.assertTrue(report['buildReport'].startswith('outputs/internal/'))
        self.assertRaises(ValueError,build,{**plan,'slides':[{**slides[0],'templateKey':'missing'}]},self.root)
    def test_prepared_output_symlink_is_rejected(self):
        self.source('source.txt',b'Unverified synthetic evidence.','text/plain');external=self.root/'untouched.txt';external.write_text('must stay intact');(self.root/'prepared').mkdir();(self.root/'prepared/evidence.json').symlink_to(external)
        self.assertRaises(ValueError,prepare,self.root);self.assertEqual(external.read_text(),'must stay intact')
    def test_supplied_architecture_image_preserves_bytes_without_inventing_text(self):
        data=(REPO/'AI_Solutioning_Agent_BRD_v1.0/02_Templates/Architecture formats/Solution_Architecture_Format.png').read_bytes();self.source('architecture.png',data,'image/png');self.assertEqual(prepare(self.root)['status'],'prepared');record=json.loads((self.root/'prepared/evidence.json').read_text())['evidence'][0];self.assertEqual(record['text'],'');self.assertEqual(record['extraction'],'visual_source');self.assertFalse(record['verified']);self.assertTrue(record['needsVisualVerification']);self.assertEqual((self.root/'prepared'/record['visual']).read_bytes(),data)
    def test_rfi_buyer_tracker_provenance_anonymity_and_literal_cells(self):
        self.source('neutral-rfp.txt',b'Section 4: The source system is not specified.\nSection 5: No price is approved.','text/plain');prepare(self.root);evidence=json.loads((self.root/'prepared/evidence.json').read_text())['evidence']
        question={'id':'Q1','rank':1,'simple':'=2+3','complex':'Which system provides our input data?','answerState':'not_found','impacts':['design'],'impactReason':'The source system changes the interface design.','searchedEvidenceIds':[e['id'] for e in evidence],'citations':[{'evidenceId':evidence[0]['id'],'quote':evidence[0]['text']}],'closestAnswerEvidenceId':evidence[1]['id']}
        plan={'questionCap':2,'anonymised':True,'buyerTemplateRequired':False,'questions':[question],'dropped':[{'question':'What is our price?','disposition':'irrelevant','reason':'No price is being agreed during this synthetic fixture.'}]}
        report=build_rfi(plan,self.root);self.assertEqual([o['visibility'] for o in report['outputs']],['client','internal'])
        with ZipFile(self.root/report['outputs'][0]['path']) as archive:
            strings=archive.read('xl/sharedStrings.xml').decode();sheet=archive.read('xl/worksheets/sheet1.xml').decode();self.assertIn('=2+3',strings);self.assertNotIn('<f>',sheet);self.assertNotIn('Closest answering',strings);self.assertNotIn('Mindmap',strings)
        with ZipFile(self.root/report['outputs'][1]['path']) as archive:self.assertIn('Closest answering passage',archive.read('xl/sharedStrings.xml').decode())
        self.source('questions.xlsx',(self.root/report['outputs'][0]['path']).read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');self.assertEqual(prepare(self.root)['status'],'prepared');self.assertTrue(any(r['kind']=='table_cell' for r in json.loads((self.root/'prepared/evidence.json').read_text())['evidence']))
    def test_rfi_rejects_unconfirmed_caps_answers_duplicates_and_fabricated_citations(self):
        self.source('neutral-source.txt',b'A synthetic clause without a defined interface.','text/plain');prepare(self.root);evidence=json.loads((self.root/'prepared/evidence.json').read_text())['evidence'][0]
        question={'id':'Q1','rank':1,'simple':'Which system supplies our data?','complex':'Please identify the source system.','answerState':'not_found','impacts':['design'],'impactReason':'The interface affects our design.','searchedEvidenceIds':[evidence['id']],'citations':[{'evidenceId':evidence['id'],'quote':evidence['text']}]}
        plan={'questionCap':2,'anonymised':True,'buyerTemplateRequired':False,'questions':[question]}
        for invalid in [{**plan,'questionCap':None},{**plan,'buyerTemplateRequired':True},{**plan,'questions':[{**question,'answerState':'found'}]},{**plan,'questions':[{**question,'simple':'Can Mindmap Digital access the system?'}]},{**plan,'questions':[{**question,'citations':[{'evidenceId':evidence['id'],'quote':'This was never stated.'}]}]},{**plan,'questions':[question,{**question,'id':'Q2','rank':2}]}]: self.assertRaises(ValueError,build_rfi,invalid,self.root)

if __name__=='__main__': unittest.main()
